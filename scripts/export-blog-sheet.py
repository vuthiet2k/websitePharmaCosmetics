# -*- coding: utf-8 -*-
"""Xuất sheet QUẢN LÝ BLOG (danh mục + bài viết + quy ước tag) cho đội nội dung.

Chạy:  python scripts/export-blog-sheet.py [đường_dẫn_ra.xlsx]   (hoặc npm run sheet:blog)
Mặc định ra: exports/quan-ly-blog.xlsx -> tải lên Google Drive, mở bằng Google Sheets.

Lý do (2026-10-13): theme đọc nhiều thứ từ tag bài viết dạng khoá:giá trị (youtube:, thoiluong:, chude:, noibat…) và
từ các blog được cấu hình trỏ tới (news_hub_blogs, video_blog_handle, expert_blog_handle…). Đội nội dung cần 1 bảng
để lên kế hoạch bài, biết blog nào đang được theme dùng ở đâu, và có sẵn chuỗi tag đúng cú pháp để dán vào Sapo.
Nguồn: configs/settings_schema.json + settings_data.json (blog nào được trỏ tới) và data/articles.js (dữ liệu demo
của bản preview — thay bằng bài thật trên Sapo khi quản lý). Không tự sinh số liệu.
"""
import json, os, subprocess, sys
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = sys.argv[1] if len(sys.argv) > 1 else os.path.join(ROOT, 'exports', 'quan-ly-blog.xlsx')

schema = json.load(open(os.path.join(ROOT, 'configs/settings_schema.json'), encoding='utf-8'))
current = json.load(open(os.path.join(ROOT, 'configs/settings_data.json'), encoding='utf-8')).get('current', {})
labels = {f['id']: (sec['name'], f.get('label', '')) for sec in schema for f in sec.get('settings', []) if f.get('id')}
setting = lambda k, d='': (current.get(k) if current.get(k) not in (None, '') else d)

# data/articles.js là module Node — đọc qua node để không chép lại logic dựng dữ liệu.
JS = ("const {blogs}=require('./data/articles');"
      "console.log(JSON.stringify(Object.entries(blogs).map(([h,b])=>({handle:h,name:b.name||b.title,url:b.url,"
      "suffix:b.template_suffix||'',articles:(b.articles||[]).map(a=>({title:a.title,handle:a.handle,url:a.url,"
      "date:a.published_on,author:a.author||'',suffix:a.template_suffix||'',tags:a.tags||[],"
      "summary:(a.summary||a.excerpt||'')}))}))))")
blogs = json.loads(subprocess.run(['node', '-e', JS], cwd=ROOT, capture_output=True, text=True, encoding='utf-8', check=True).stdout)

FEATURED_TAG = setting('news_featured_tag', 'noibat')
HUB = [h.strip() for h in str(setting('news_hub_blogs', '')).split(',') if h.strip()]

# --- Style (cùng bộ với scripts/export-config-sheet.py) ---
GREEN = '1F6F54'
F_HEAD = Font(bold=True, color='FFFFFF'); P_HEAD = PatternFill('solid', fgColor=GREEN)
P_GROUP = PatternFill('solid', fgColor='DDEFE6'); P_EDIT = PatternFill('solid', fgColor='FFF7D6')
P_AUTO = PatternFill('solid', fgColor='EEF3F0')
THIN = Side(style='thin', color='D0D0D0'); BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical='top')
STATUS = ['Chưa xử lý', 'Đang viết', 'Chờ duyệt', 'Đã đăng', 'Cần sửa', 'Gỡ bài']
ACTIONS = ['Giữ nguyên', 'Đăng mới', 'Sửa nội dung', 'Sửa tag', 'Đổi danh mục', 'Gỡ bài']

def header_row(ws, cols, widths, freeze_col=3):
    ws.append(cols)
    for c, w in zip(ws[ws.max_row], widths):
        c.font, c.fill, c.alignment, c.border = F_HEAD, P_HEAD, Alignment(wrap_text=True, vertical='center'), BORDER
        ws.column_dimensions[c.column_letter].width = w
    ws.row_dimensions[ws.max_row].height = 34
    # Dùng toạ độ chữ: ws.cell(...) sẽ tạo sẵn ô ở dòng dưới ⇒ lần append sau bị đẩy xuống, chừa 1 dòng trống.
    ws.freeze_panes = f'{get_column_letter(freeze_col)}{ws.max_row + 1}'

def style_row(ws, r, edit_cols=(), auto_cols=()):
    for c in ws[r]: c.alignment, c.border = WRAP, BORDER
    for col in edit_cols: ws.cell(r, col).fill = P_EDIT
    for col in auto_cols: ws.cell(r, col).fill = P_AUTO

def split_tags(tags):
    """Tách tag theo quy ước theme → các cột của sheet Bài viết."""
    out = {'youtube': [], 'facebook': [], 'thoiluong': '', 'chude': '', 'chuyengia': '', 'noibat': '', 'khac': []}
    for t in tags:
        key, _, val = t.partition(':')
        k = key.strip().lower()
        if val and k in ('youtube', 'yt'): out['youtube'].append(val.strip())
        elif val and k in ('facebook', 'fb'): out['facebook'].append(val.strip())
        elif val and k == 'thoiluong': out['thoiluong'] = val.strip()
        elif val and k == 'chude': out['chude'] = val.strip()
        elif val and k == 'chuyengia': out['chuyengia'] = val.strip()
        elif t.strip() == FEATURED_TAG: out['noibat'] = 'Có'
        else: out['khac'].append(t)
    return out

wb = Workbook()

# ===== Hướng dẫn =====
g = wb.active; g.title = 'Hướng dẫn'
g.column_dimensions['A'].width = 26; g.column_dimensions['B'].width = 46; g.column_dimensions['C'].width = 70
def gl(*vals, bold=False, fill=None):
    g.append(list(vals))
    for c in g[g.max_row]:
        c.alignment = WRAP
        if bold: c.font = Font(bold=True)
        if fill: c.fill = fill
gl('QUẢN LÝ BLOG — PHARMA COSMETICS'); g['A1'].font = Font(bold=True, size=15, color=GREEN)
gl('Nguồn', 'Sinh tự động bằng: npm run sheet:blog. Danh mục + vai trò lấy từ cấu hình theme; bài viết lấy từ dữ liệu demo của bản preview — khi quản lý thật, cập nhật theo bài trên Sapo.')
gl()
gl('CÁCH DÙNG', bold=True, fill=P_GROUP)
for a, b in [
    ('Sheet "Danh mục"', 'Mỗi dòng = 1 danh mục bài viết (blog) trên Sapo: đường dẫn, giao diện phải chọn, theme đang dùng danh mục ở đâu. Không đổi Handle của danh mục đang được cấu hình trỏ tới (cột Cấu hình đang dùng) — đổi là khối đó trên web trống.'),
    ('Sheet "Bài viết"', 'Mỗi dòng = 1 bài. Muốn đăng bài mới: thêm dòng, điền các cột NỀN VÀNG. Cột "Tag dán vào Sapo" (nền xanh nhạt) TỰ GHÉP từ các cột tag — chép nguyên ô đó dán vào ô Tag của bài trên Sapo.'),
    ('Giao diện', 'Bài trong danh mục Video chọn giao diện article.video; danh mục Video chọn blog.video. Các bài/danh mục khác để giao diện mặc định.'),
    ('Trang tổng tin tức', 'Trang /tin-tuc hiện các danh mục ở cấu hình "Danh mục hiển thị" theo thứ tự: ' + (', '.join(HUB) or '(chưa cấu hình)') + '. Danh mục đầu lên khối "Tin mới nhất".'),
    ('Quy tắc nội dung', 'Dùng "Chuyên gia", không dùng "Bác sĩ". Không bịa số liệu (đánh giá, % hiệu quả, lượt xem). Video phải có thật, công khai.'),
]: gl(a, b)
gl()
gl('QUY ƯỚC TAG THEME ĐANG ĐỌC', bold=True, fill=P_GROUP)
gl('Tag', 'Ví dụ', 'Tác dụng trên web', bold=True)
for row in [
    ('youtube:<mã hoặc link>', 'youtube:4B9RN_wcl_I', 'Nguồn video YouTube (mã 11 ký tự, phân biệt hoa/thường). Bài thường: nhiều tag = video chính + video bổ trợ.'),
    ('facebook:fb_video_<số video>', 'facebook:fb_video_427805549152968', 'Nguồn video Facebook (video Công khai trên Fanpage). Số video = dãy số sau /videos/, ?v= hoặc /reel/ trên thanh địa chỉ khi mở video. Ghi facebook:<số> cũng được. Bắt buộc có ảnh đại diện.'),
    ('khung:doc', 'khung:doc', 'Video quay dọc (9:16) ⇒ khung phát dọc, không bị thu nhỏ. Link /reel/ tự nhận, không cần tag. Ghi ở cột "Tag khác".'),
    ('thoiluong:<MM:SS>', 'thoiluong:01:23', 'Thời lượng video trên ảnh bìa + dữ liệu SEO.'),
    ('chude:<Nhãn>', 'chude:Làm sạch da', 'Nhãn chủ đề trên thẻ bài/video và nút lọc trang Video. Gõ giống hệt nhau giữa các bài.'),
    ('chuyengia:<handle>', 'chuyengia:chuyen-gia-thanh-hang', 'Hiện ảnh + tên chuyên gia (bài trong danh mục "' + setting('expert_blog_handle', 'chuyen-gia') + '").'),
    (FEATURED_TAG, FEATURED_TAG, 'Đưa bài lên cột "Tin nổi bật" ở trang tổng /tin-tuc.'),
    ('sanpham:<alias sản phẩm>', 'sanpham:serum-vitamin-c-15', 'Gắn bài với sản phẩm: "Bài viết liên quan", đánh giá, FAQ, ý kiến chuyên gia ở trang sản phẩm.'),
    ('hoatchat:<handle>', 'hoatchat:vitamin-c', 'Ghép bài với sản phẩm cùng hoạt chất ở "Bài viết liên quan".'),
]: gl(*row)

# ===== Danh mục =====
used_by = {}
for k, v in current.items():
    if isinstance(v, str) and v.strip():
        for h in [x.strip() for x in v.split(',')]:
            used_by.setdefault(h, []).append(k)
dm = wb.create_sheet('Danh mục')
DM_COLS = ['STT', 'Handle (đường dẫn)', 'Tên danh mục', 'URL', 'Giao diện cần chọn', 'Số bài (demo)', 'Trang tổng /tin-tuc',
           'Cấu hình đang dùng (theme hiện ở đâu)', 'Hành động ★', 'Đề xuất / Ghi chú ★', 'Trạng thái ★']
header_row(dm, DM_COLS, [6, 24, 28, 30, 18, 10, 16, 60, 16, 40, 14])
dv_s1 = DataValidation(type='list', formula1='"' + ','.join(STATUS) + '"', allow_blank=True)
dv_a1 = DataValidation(type='list', formula1='"Giữ nguyên,Tạo mới,Đổi tên,Gỡ"', allow_blank=True)
dm.add_data_validation(dv_s1); dm.add_data_validation(dv_a1)
for i, b in enumerate(blogs, 1):
    refs = used_by.get(b['handle'], [])
    where = '\n'.join(f"{labels.get(k, ('', k))[0]} › {labels.get(k, ('', k))[1]}  [{k}]" for k in refs)
    hub = f"Có (thứ {HUB.index(b['handle']) + 1})" if b['handle'] in HUB else ''
    dm.append([i, b['handle'], b['name'], b['url'], 'blog.' + b['suffix'] if b['suffix'] else 'blog (mặc định)',
               len(b['articles']), hub, where or '(chưa cấu hình nơi nào dùng)', 'Giữ nguyên', '', 'Đã đăng'])
    r = dm.max_row
    style_row(dm, r, edit_cols=(9, 10, 11))
    dm.cell(r, 2).font = Font(name='Consolas', size=10, color='7A1F1F')
    dv_a1.add(dm.cell(r, 9)); dv_s1.add(dm.cell(r, 11))
for r in range(dm.max_row + 1, dm.max_row + 11):
    dm.append([r - 1]); style_row(dm, r, edit_cols=(2, 3, 5, 9, 10, 11)); dv_a1.add(dm.cell(r, 9)); dv_s1.add(dm.cell(r, 11))
n_dm = dm.max_row

# ===== Bài viết =====
bv = wb.create_sheet('Bài viết')
BV_COLS = ['STT', 'Danh mục (handle) ★', 'Tiêu đề ★', 'Handle bài', 'Ngày đăng', 'Tác giả', 'Giao diện ★',
           'YouTube (mã/link) ★', 'Facebook (số video) ★', 'Thời lượng MM:SS ★', 'Chủ đề ★', 'Nổi bật ★', 'Chuyên gia (handle) ★',
           'Tag khác ★ (phẩy)', 'Tag dán vào Sapo (tự ghép)', 'Tóm tắt', 'Hành động ★', 'Ghi chú ★', 'Trạng thái ★']
header_row(bv, BV_COLS, [6, 20, 46, 30, 12, 20, 15, 18, 30, 11, 20, 9, 22, 26, 52, 46, 14, 28, 13], freeze_col=4)
bv.auto_filter.ref = f'A1:{get_column_letter(len(BV_COLS))}1'
EDIT = (2, 3, 7, 8, 9, 10, 11, 12, 13, 14, 17, 18, 19)
dv_blog = DataValidation(type='list', formula1=f"='Danh mục'!$B$2:$B${n_dm}", allow_blank=True)
dv_tpl = DataValidation(type='list', formula1='"article,article.video"', allow_blank=True)
dv_yes = DataValidation(type='list', formula1='"Có,Không"', allow_blank=True)
dv_act = DataValidation(type='list', formula1='"' + ','.join(ACTIONS) + '"', allow_blank=True)
dv_st = DataValidation(type='list', formula1='"' + ','.join(STATUS) + '"', allow_blank=True)
for dv in (dv_blog, dv_tpl, dv_yes, dv_act, dv_st): bv.add_data_validation(dv)

def tag_formula(r):
    # Ghép tag đúng cú pháp theme; YouTube/Facebook có thể nhiều giá trị cách nhau dấu phẩy (video chính trước).
    yt = f'IF(H{r}="","","youtube:"&SUBSTITUTE(SUBSTITUTE(H{r}," ",""),",",", youtube:"))'
    fb = f'IF(I{r}="","","facebook:"&SUBSTITUTE(SUBSTITUTE(I{r}," ",""),",",", facebook:"))'
    return (f'=_xlfn.TEXTJOIN(", ",TRUE,{yt},{fb},IF(J{r}="","","thoiluong:"&J{r}),IF(K{r}="","","chude:"&K{r}),'
            f'IF(L{r}="Có","{FEATURED_TAG}",""),IF(M{r}="","","chuyengia:"&M{r}),N{r})')

n = 0
for b in blogs:
    for a in b['articles']:
        n += 1; t = split_tags(a['tags'])
        bv.append([n, b['handle'], a['title'], a['handle'], (a['date'] or '')[:10], a['author'],
                   'article.' + a['suffix'] if a['suffix'] else 'article',
                   ', '.join(t['youtube']), ', '.join(t['facebook']), t['thoiluong'], t['chude'], t['noibat'] or 'Không',
                   t['chuyengia'], ', '.join(t['khac']), None, a['summary'], 'Giữ nguyên', '', 'Đã đăng'])
        r = bv.max_row
        bv.cell(r, 15).value = tag_formula(r)
        style_row(bv, r, edit_cols=EDIT, auto_cols=(15,))
        bv.cell(r, 15).font = Font(name='Consolas', size=9)
        for col, dv in ((2, dv_blog), (7, dv_tpl), (12, dv_yes), (17, dv_act), (19, dv_st)): dv.add(bv.cell(r, col))
for _ in range(30):  # dòng trống để lên kế hoạch bài mới
    n += 1; bv.append([n]); r = bv.max_row
    bv.cell(r, 15).value = tag_formula(r)
    style_row(bv, r, edit_cols=EDIT, auto_cols=(15,))
    bv.cell(r, 15).font = Font(name='Consolas', size=9)
    for col, dv in ((2, dv_blog), (7, dv_tpl), (12, dv_yes), (17, dv_act), (19, dv_st)): dv.add(bv.cell(r, col))

os.makedirs(os.path.dirname(OUT), exist_ok=True)
wb.save(OUT)
print(f'OK: {os.path.relpath(OUT, ROOT)} — {len(blogs)} danh mục, {sum(len(b["articles"]) for b in blogs)} bài (+30 dòng trống)')
